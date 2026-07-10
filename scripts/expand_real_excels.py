# -*- coding: utf-8 -*-
"""扩充真实双 Excel（主表 + 面试表）到指定数据行数，供导入压测与手工验证。

合并自原 expand_real_excels_10000.py 与 expand_real_excels_50000.py，
两者仅目标路径/行数/写入方式不同，字段构造逻辑完全一致。

用法：
    python scripts/expand_real_excels.py                          # 根目录双表，upsert 到 10000 行
    python scripts/expand_real_excels.py --target fixtures --rows 50000 --rewrite
                                                                  # fixtures 双表，流式重写 50000 行

写入方式：
- 默认 upsert：按「应聘档案编号」定位已有行覆盖，保留原始真实数据行，可重复运行。
- --rewrite：write_only 流式全量重写（只保留表头），5 万行级别内存友好，适合压测夹具。

生成数据被 tests/test_master_import_parse.py 与 scripts/benchmark_master_import.py 消费；
阶段状态组合与 campus/domain/stage_routing.py 的判定规则对应（见 config/master_import/registration/stage_rules.json）。
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures" / "master_import"
APP_NAME = "applicationProcessList20260626083908.xlsx"
INTERVIEW_NAME = "候选人面试安排管理列表20260626083900.xlsx"

# 每个元组对应一种"候选人停留在某阶段"的真实状态组合，顺序:
# (标签, 登记, 简历筛选, 资审, 商密, 笔试, 当前环节, 技术面, 主管面, 报批, 谈薪)
STAGES = [
    ("投递", "待投递", "", "", "", "", "", "", "", "", ""),
    ("简历筛选", "已登记", "待筛选", "", "", "", "", "", "", "", ""),
    ("资格审查", "已登记", "通过", "待审查", "待签署", "", "资审", "", "", "", ""),
    ("商业秘密签署", "已登记", "通过", "通过", "待签署", "待预约", "商业秘密签署", "", "", "", ""),
    ("笔试", "已登记", "通过", "通过", "已签署", "已预约", "笔试", "", "", "", ""),
    ("性格测评", "已登记", "通过", "通过", "已签署", "已完成", "性格测评", "", "", "", ""),
    ("资格面试", "已登记", "通过", "通过", "已签署", "已完成", "资格面试", "", "", "", ""),
    ("技术面", "已登记", "通过", "通过", "已签署", "已完成", "技术面", "已预约", "", "", ""),
    ("主管面", "已登记", "通过", "通过", "已签署", "已完成", "主管面", "已完成", "已预约", "", ""),
    ("offer策略", "已登记", "通过", "通过", "已签署", "已完成", "报批", "已完成", "已完成", "审批中", "待谈薪"),
]

_LATE_STAGES = ("技术面", "主管面", "offer策略")


def archive_id(seq: int) -> str:
    return f"SR20260709{seq:05d}"


def case_for(seq: int) -> dict:
    label, reg, resume, qual, secret, written, step, tech, manager, approval, salary = \
        STAGES[(seq - 1) % len(STAGES)]
    return {
        "seq": seq,
        "aid": archive_id(seq),
        "label": label,
        "name": f"扩展候选人{seq:05d}-{label}",
        "phone": f"1378{seq:07d}"[-11:],
        "registration": reg,
        "resume": resume,
        "qualification": qual,
        "secret": secret,
        "written": written,
        "step": step,
        "tech": tech,
        "manager": manager,
        "approval": approval,
        "salary": salary,
    }


# 字段规格：(候选表头名列表, 取值函数)。表头别名与
# config/master_import/registration/field_mappings.json 的列匹配保持一致。
APPLICATION_SPEC = [
    (("应聘档案编号",), lambda c: c["aid"]),
    (("简历编号",), lambda c: ""),  # 业务上已弃用，真实表保留历史列则留空
    (("姓名", "候选人", "候选人姓名"), lambda c: c["name"]),
    (("联系电话", "电话", "候选人手机号"), lambda c: c["phone"]),
    (("学历", "学历（1）"), lambda c: "硕士"),
    (("毕业学校（1）", "毕业院校", "学校"), lambda c: "测试大学"),
    (("专业（1）", "专业"), lambda c: "软件工程"),
    (("登记状态", "应聘状态"), lambda c: c["registration"]),
    (("简历筛选状态",), lambda c: c["resume"]),
    (("资审状态",), lambda c: c["qualification"]),
    (("商业秘密签署状态", "商业合规签署通知状态", "保密协议状态"), lambda c: c["secret"]),
    (("笔试状态", "考试状态"), lambda c: c["written"]),
    (("当前环节",), lambda c: c["step"]),
    (("当前环节状态",), lambda c: "进行中" if c["step"] else ""),
    (("技术面状态", "专业面试状态"), lambda c: c["tech"]),
    (("主管面状态", "业务主管面试状态"), lambda c: c["manager"]),
    (("报批状态",), lambda c: c["approval"]),
    (("谈薪状态",), lambda c: c["salary"]),
    (("Offer状态",), lambda c: "未发放" if c["label"] == "offer策略" else ""),
    (("签约状态",), lambda c: "未签约"),
    (("是否入职",), lambda c: "否"),
    (("二层部门", "录用二级部门"), lambda c: "软件部"),
    (("三层部门", "录用三级部门"), lambda c: "研发一组"),
    (("拓源人",), lambda c: "hr01"),
    (("拓源人信息", "拓源人部门"), lambda c: "软件部"),
    (("接口人",), lambda c: "hr02"),
    (("接口人信息", "接口人部门"), lambda c: "软件部"),
    (("备注说明",), lambda c: f"{c['label']}流程扩展测试"),
]

INTERVIEW_SPEC = [
    (("应聘档案编号",), lambda c: c["aid"]),
    (("候选人编号",), lambda c: c["aid"]),
    (("候选人姓名",), lambda c: c["name"]),
    (("候选人手机号",), lambda c: "+86 XXXXXXXXXXX"),
    (("当前环节",), lambda c: c["step"]),
    (("当前环节状态",), lambda c: "进行中" if c["step"] else ""),
    (("面试进展",), lambda c: c["step"]),
    (("面试考核状态",), lambda c: "考核中" if c["seq"] % 10 >= 6 else ""),
    (("综测状态",), lambda c: "通过" if c["label"] in ("性格测评", "资格面试") + _LATE_STAGES else ""),
    (("考试状态",), lambda c: c["written"]),
    (("专业面试-面试结论", "专业面试1-面试结论"), lambda c: "A" if c["label"] in _LATE_STAGES else ""),
    (("业务主管面试-面试结论",), lambda c: "通过" if c["label"] in ("主管面", "offer策略") else ""),
    (("面试结论",), lambda c: "A" if c["label"] in _LATE_STAGES else ""),
]


def header_index(headers: list) -> dict:
    return {str(h).strip(): i for i, h in enumerate(headers) if h is not None and str(h).strip()}


def read_headers(path: Path) -> list:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        row = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
        return [c if c is not None else "" for c in row]
    finally:
        wb.close()


def expand_upsert(path: Path, spec, rows: int, template: Path | None = None) -> None:
    """按应聘档案编号原地 upsert，保留文件中已有的真实数据行。"""
    src = path if path.exists() else template
    if src is None or not Path(src).exists():
        raise FileNotFoundError(f"目标文件不存在且无可用模板: {path}")
    wb = load_workbook(src)
    ws = wb.active
    hm = {str(c.value).strip(): idx + 1 for idx, c in enumerate(ws[1])
          if c.value is not None and str(c.value).strip()}
    if "应聘档案编号" not in hm:
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value="应聘档案编号")
        hm["应聘档案编号"] = col
    key_col = hm["应聘档案编号"]
    row_map = {}
    for row in range(2, ws.max_row + 1):
        aid = str(ws.cell(row=row, column=key_col).value or "").strip()
        if aid:
            row_map[aid] = row
    for seq in range(1, rows + 1):
        case = case_for(seq)
        row = row_map.get(case["aid"]) or ws.max_row + 1
        for names, value_fn in spec:
            for name in names:
                if name in hm:
                    ws.cell(row=row, column=hm[name], value=value_fn(case))
        row_map[case["aid"]] = row
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True).active
    print(f"{path.name}: data_rows={check.max_row - 1}, cols={check.max_column} (upsert)")


def expand_rewrite(path: Path, spec, rows: int, sheet_title: str) -> None:
    """write_only 流式全量重写（仅保留表头），5 万行级别内存友好。"""
    if not path.exists():
        raise FileNotFoundError(f"模板不存在: {path}")
    headers = read_headers(path)
    hm = header_index(headers)
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title)
    ws.append(headers)
    for seq in range(1, rows + 1):
        case = case_for(seq)
        row = [""] * len(headers)
        for names, value_fn in spec:
            for name in names:
                if name in hm:
                    row[hm[name]] = value_fn(case)
        ws.append(row)
        if seq % 10000 == 0:
            print(f"  … {path.name}: {seq}/{rows}")
    wb.save(tmp)
    tmp.replace(path)
    print(f"{path.name}: data_rows={rows}, cols={len(headers)}, sheet={sheet_title} (rewrite)")


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--target", choices=("root", "fixtures"), default="root",
                        help="root=仓库根目录双表（手工导入验证）；fixtures=tests/fixtures 夹具（压测）")
    parser.add_argument("--rows", type=int, default=10000, help="目标数据行数")
    parser.add_argument("--rewrite", action="store_true",
                        help="流式全量重写（丢弃原有数据行）；不加则原地 upsert 保留已有行")
    args = parser.parse_args()

    base = ROOT if args.target == "root" else FIXTURES
    app_path = base / APP_NAME
    interview_path = base / INTERVIEW_NAME

    if args.rewrite:
        expand_rewrite(app_path, APPLICATION_SPEC, args.rows, "Sheet")
        expand_rewrite(interview_path, INTERVIEW_SPEC, args.rows, "Sheet1")
    else:
        expand_upsert(app_path, APPLICATION_SPEC, args.rows)
        expand_upsert(interview_path, INTERVIEW_SPEC, args.rows,
                      template=FIXTURES / INTERVIEW_NAME)


if __name__ == "__main__":
    main()
