# -*- coding: utf-8 -*-
"""向根目录真实主数据双 Excel 写入流程状态测试样例。

根目录双表由 tests/fixtures/master_import 下的真实业务表复制而来（列结构
与真实表完全一致，不新增列），再按「应聘档案编号」upsert 样例行，重复运行
不会重复追加。阶段判定完全依赖真实表中的「当前环节」（应聘表列名为
“当前缓解”）：覆盖 01-投递 到 10-offer策略。
"""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
APP_FILE = ROOT / "applicationProcessList20260626083908.xlsx"
INTERVIEW_FILE = ROOT / "候选人面试安排管理列表20260626083900.xlsx"
FIXTURE_DIR = ROOT / "tests" / "fixtures" / "master_import"


def sr_id(date_str, seq):
    return f"SR{date_str}{seq:05d}"


#: 每个样例由「当前环节」唯一驱动阶段（真实表中不存在各流程状态列）。
CASES = [
    {
        "seq": 1, "date": "20260301", "name": "真实状态01投递", "phone": "13790003001",
        "current_step": "",
        "expected_stage": "registration", "expected_status": "01-投递",
    },
    {
        "seq": 2, "date": "20260302", "name": "真实状态02简历筛选", "phone": "13790003002",
        "current_step": "简历筛选",
        "expected_stage": "resume_screening", "expected_status": "02-简历筛选",
    },
    {
        "seq": 3, "date": "20260303", "name": "真实状态03资格审查", "phone": "13790003003",
        "current_step": "资审",
        "expected_stage": "qualification", "expected_status": "03-资格审查",
    },
    {
        "seq": 4, "date": "20260304", "name": "真实状态04商业秘密签署", "phone": "13790003004",
        "current_step": "商业秘密签署",
        "expected_stage": "commercial_secret", "expected_status": "04-商业秘密签署",
    },
    {
        "seq": 5, "date": "20260305", "name": "真实状态05笔试", "phone": "13790003005",
        "current_step": "笔试",
        "expected_stage": "written_test", "expected_status": "05-笔试",
    },
    {
        "seq": 6, "date": "20260306", "name": "真实状态06性格测评", "phone": "13790003006",
        "current_step": "性格测评",
        "expected_stage": "personality_test", "expected_status": "06-性格测评",
    },
    {
        "seq": 7, "date": "20260307", "name": "真实状态07资格面试", "phone": "13790003007",
        "current_step": "资格面试",
        "expected_stage": "qualification_interview", "expected_status": "07-资格面试",
    },
    {
        "seq": 8, "date": "20260308", "name": "真实状态08技术面", "phone": "13790003008",
        "current_step": "技术面",
        "expected_stage": "tech_interview", "expected_status": "08-技术面",
    },
    {
        "seq": 9, "date": "20260309", "name": "真实状态09主管面", "phone": "13790003009",
        "current_step": "主管面",
        "expected_stage": "manager_interview", "expected_status": "09-主管面",
    },
    {
        "seq": 10, "date": "20260310", "name": "真实状态10offer策略", "phone": "13790003010",
        "current_step": "报批",
        "expected_stage": "approval", "expected_status": "10-offer策略",
    },
]


def col_index(ws, header):
    """真实表既有列的序号；不存在的列直接报错（不允许新增列）。"""
    headers = [str(c.value or "").strip() for c in ws[1]]
    if header not in headers:
        raise ValueError(f"真实表中不存在列「{header}」")
    return headers.index(header) + 1


def upsert_row(ws, archive_id, values):
    key_col = col_index(ws, "应聘档案编号")
    row_idx = None
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=key_col).value or "").strip() == archive_id:
            row_idx = row
            break
    if row_idx is None:
        row_idx = ws.max_row + 1
    for header, value in values.items():
        ws.cell(row=row_idx, column=col_index(ws, header), value=value)
    return row_idx


def app_values(case):
    """applicationProcessList 真实列（当前环节列名为“当前缓解”）。"""
    return {
        "应聘档案编号": sr_id(case["date"], case["seq"]),
        "姓名": case["name"],
        "联系电话": case["phone"],
        "学历（1）": "硕士",
        "毕业学校（1）": "测试大学",
        "专业（1）": "软件工程",
        "应聘状态": "进行中",
        "当前缓解": case["current_step"],
        "当前环节状态": "进行中" if case["current_step"] else "",
    }


def interview_values(case):
    """候选人面试安排管理列表真实列。"""
    rid = sr_id(case["date"], case["seq"])
    return {
        "应聘档案编号": rid,
        "候选人编号": rid,
        "候选人姓名": case["name"],
        "候选人手机号": "+86 XXXXXXXXXXX",
        "当前环节": case["current_step"],
        "面试进展": case["current_step"],
        "考试状态": "已完成" if case["seq"] >= 5 else "",
        "综测状态": "通过" if case["seq"] >= 6 else "",
        "专业面试-面试结论": "A" if case["seq"] >= 8 else "",
        "业务主管面试-面试结论": "通过" if case["seq"] >= 9 else "",
        "面试结论": "A" if case["seq"] >= 8 else "",
    }


def _sync_from_fixture(target: Path):
    """根目录副本缺失或列结构与真实表不一致时，从 fixtures 重新复制。"""
    src = FIXTURE_DIR / target.name
    if not src.is_file():
        raise FileNotFoundError(f"真实业务表不存在：{src}")
    if target.is_file():
        src_wb = load_workbook(src, read_only=True)
        dst_wb = load_workbook(target, read_only=True)
        same = [str(c or "").strip() for c in next(src_wb.active.iter_rows(max_row=1, values_only=True))] \
            == [str(c or "").strip() for c in next(dst_wb.active.iter_rows(max_row=1, values_only=True))]
        src_wb.close()
        dst_wb.close()
        if same:
            return
    shutil.copyfile(src, target)


def expand():
    _sync_from_fixture(APP_FILE)
    _sync_from_fixture(INTERVIEW_FILE)
    app_wb = load_workbook(APP_FILE)
    app_ws = app_wb.active
    interview_wb = load_workbook(INTERVIEW_FILE)
    interview_ws = interview_wb.active
    for case in CASES:
        rid = sr_id(case["date"], case["seq"])
        upsert_row(app_ws, rid, app_values(case))
        # 只给确实需要面试/环节表推动的流程补面试表，避免早期阶段被面试进展误命中。
        if case["seq"] >= 5:
            upsert_row(interview_ws, rid, interview_values(case))
    app_wb.save(APP_FILE)
    interview_wb.save(INTERVIEW_FILE)
    return len(CASES)


if __name__ == "__main__":
    print(f"已写入/更新真实流程状态样例 {expand()} 条")
