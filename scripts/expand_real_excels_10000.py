# -*- coding: utf-8 -*-
"""扩充根目录真实双 Excel 到至少 10000 条数据行。

目标文件：
- applicationProcessList20260626083908.xlsx
- 候选人面试安排管理列表20260626083900.xlsx

若根目录面试表不存在，则以 tests/fixtures/master_import 下同名样例为模板创建。
按「应聘档案编号」upsert 固定 10000 条扩展候选人，重复运行不会重复追加。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "applicationProcessList20260626083908.xlsx"
INTERVIEW = ROOT / "候选人面试安排管理列表20260626083900.xlsx"
INTERVIEW_TEMPLATE = ROOT / "tests" / "fixtures" / "master_import" / "候选人面试安排管理列表20260626083900.xlsx"
TARGET_ROWS = 10000

STAGES = [
    ("投递", "待投递", "", "", "", "", "", "", "", "", ""),
    ("简历筛选", "已登记", "待筛选", "", "", "", "", "", "", "", ""),
    ("资格审查", "已登记", "通过", "待审查", "待签署", "", "资审", "", "", "", ""),
    ("商业秘密签署", "已登记", "通过", "通过", "待签署", "待预约", "商业秘密签署", "", "", "", ""),
    ("笔试", "已登记", "通过", "通过", "已签署", "已预约", "笔试", "", "", "", ""),
    ("性格测评", "已登记", "通过", "通过", "已签署", "已完成", "性格测评", "", "", "", ""),
    ("资格面试", "已登记", "通过", "通过", "已签署", "已完成", "资格面试", "", "", "", ""),
    ("技术面", "已登记", "通过", "通过", "已签署", "已完成", "技术面", "已预约", "", "", ""),
    ("主管面", "已登记", "通过", "通过", "已签署", "已完成", "主管面", "已完成", "已预约", "", ""),
    ("offer策略", "已登记", "通过", "通过", "已签署", "已完成", "报批", "已完成", "已完成", "审批中", "待谈薪"),
]


def load_target(path: Path, template: Path | None = None):
    src = path if path.exists() else template
    if src is None or not Path(src).exists():
        raise FileNotFoundError(f"目标文件不存在且无可用模板: {path}")
    return load_workbook(src)


def headers(ws):
    return {str(c.value).strip(): idx + 1 for idx, c in enumerate(ws[1]) if c.value is not None and str(c.value).strip()}


def ensure_col(ws, hm, name):
    if name in hm:
        return hm[name]
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value=name)
    hm[name] = col
    return col


def setv(ws, row, hm, names, value):
    for name in names:
        if name in hm:
            ws.cell(row=row, column=hm[name], value=value)


def rows_by_archive(ws, hm):
    key_col = ensure_col(ws, hm, "应聘档案编号")
    out = {}
    for row in range(2, ws.max_row + 1):
        aid = str(ws.cell(row=row, column=key_col).value or "").strip()
        if aid:
            out[aid] = row
    return out


def archive_id(seq):
    return f"SR20260709{seq:05d}"


def case_for(seq):
    label, reg, resume, qual, secret, written, step, tech, manager, approval, salary = STAGES[(seq - 1) % len(STAGES)]
    aid = archive_id(seq)
    return {
        "seq": seq,
        "aid": aid,
        "label": label,
        "name": f"扩展候选人{seq:05d}-{label}",
        "phone": f"1378{seq:07d}"[-11:],
        "registration": reg,
        "resume": resume,
        "qualification": qual,
        "secret": secret,
        "written": written,
        "step": step,
        "tech": tech,
        "manager": manager,
        "approval": approval,
        "salary": salary,
    }


def fill_application(ws, row, hm, case):
    setv(ws, row, hm, ["应聘档案编号"], case["aid"])
    # 业务上不再使用“简历编号”；真实表若保留历史列则扩展行留空。
    setv(ws, row, hm, ["简历编号"], "")
    setv(ws, row, hm, ["姓名", "候选人", "候选人姓名"], case["name"])
    setv(ws, row, hm, ["联系电话", "电话", "候选人手机号"], case["phone"])
    setv(ws, row, hm, ["学历", "学历（1）"], "硕士")
    setv(ws, row, hm, ["毕业学校（1）", "毕业院校", "学校"], "测试大学")
    setv(ws, row, hm, ["专业（1）", "专业"], "软件工程")
    setv(ws, row, hm, ["登记状态", "应聘状态"], case["registration"])
    setv(ws, row, hm, ["简历筛选状态"], case["resume"])
    setv(ws, row, hm, ["资审状态"], case["qualification"])
    setv(ws, row, hm, ["商业秘密签署状态", "商业合规签署通知状态", "保密协议状态"], case["secret"])
    setv(ws, row, hm, ["笔试状态", "考试状态"], case["written"])
    setv(ws, row, hm, ["当前环节"], case["step"])
    setv(ws, row, hm, ["当前环节状态"], "进行中" if case["step"] else "")
    setv(ws, row, hm, ["技术面状态", "专业面试状态"], case["tech"])
    setv(ws, row, hm, ["主管面状态", "业务主管面试状态"], case["manager"])
    setv(ws, row, hm, ["报批状态"], case["approval"])
    setv(ws, row, hm, ["谈薪状态"], case["salary"])
    setv(ws, row, hm, ["Offer状态"], "未发放" if case["label"] == "offer策略" else "")
    setv(ws, row, hm, ["签约状态"], "未签约")
    setv(ws, row, hm, ["是否入职"], "否")
    setv(ws, row, hm, ["二层部门", "录用二级部门"], "软件部")
    setv(ws, row, hm, ["三层部门", "录用三级部门"], "研发一组")
    setv(ws, row, hm, ["拓源人"], "hr01")
    setv(ws, row, hm, ["拓源人部门"], "软件部")
    setv(ws, row, hm, ["接口人"], "hr02")
    setv(ws, row, hm, ["接口人部门"], "软件部")
    setv(ws, row, hm, ["备注说明"], f"{case['label']}流程扩展测试")


def fill_interview(ws, row, hm, case):
    setv(ws, row, hm, ["应聘档案编号"], case["aid"])
    setv(ws, row, hm, ["候选人编号"], case["aid"])
    setv(ws, row, hm, ["候选人姓名"], case["name"])
    setv(ws, row, hm, ["候选人手机号"], "+86 XXXXXXXXXXX")
    setv(ws, row, hm, ["当前环节"], case["step"])
    setv(ws, row, hm, ["当前环节状态"], "进行中" if case["step"] else "")
    setv(ws, row, hm, ["面试进展"], case["step"])
    setv(ws, row, hm, ["面试考核状态"], "考核中" if case["seq"] % 10 >= 6 else "")
    setv(ws, row, hm, ["综测状态"], "通过" if case["label"] in ("性格测评", "资格面试", "技术面", "主管面", "offer策略") else "")
    setv(ws, row, hm, ["考试状态"], case["written"])
    setv(ws, row, hm, ["专业面试-面试结论", "专业面试1-面试结论"], "A" if case["label"] in ("技术面", "主管面", "offer策略") else "")
    setv(ws, row, hm, ["业务主管面试-面试结论"], "通过" if case["label"] in ("主管面", "offer策略") else "")
    setv(ws, row, hm, ["面试结论"], "A" if case["label"] in ("技术面", "主管面", "offer策略") else "")


def expand(path: Path, filler, template: Path | None = None):
    wb = load_target(path, template)
    ws = wb.active
    hm = headers(ws)
    row_map = rows_by_archive(ws, hm)
    for seq in range(1, TARGET_ROWS + 1):
        case = case_for(seq)
        row = row_map.get(case["aid"]) or ws.max_row + 1
        filler(ws, row, hm, case)
        row_map[case["aid"]] = row
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True).active
    print(f"{path.name}: data_rows={check.max_row - 1}, cols={check.max_column}")


def main():
    expand(APP, fill_application)
    expand(INTERVIEW, fill_interview, INTERVIEW_TEMPLATE)


if __name__ == "__main__":
    main()
