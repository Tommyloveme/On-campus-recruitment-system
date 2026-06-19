# -*- coding: utf-8 -*-
"""主数据表导入与流程阶段判定引擎。"""
import fnmatch
import json
import os
from datetime import datetime

from campus.config_loader import get_stage_meta, load_stages_meta
from campus.settings import BASE_DIR

MASTER_IMPORT_DIR = os.path.join(BASE_DIR, "config", "master_import")


def load_field_mappings(page_dir, index):
    """加载网页字段与 Excel 列映射配置。"""
    fname = index.get("field_mappings_file", "field_mappings.json")
    path = os.path.join(page_dir, fname)
    if not os.path.exists(path):
        return {"fields": []}
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _mapping_columns_for_source(field_mappings, source_key):
    """将 field_mappings 中属于某数据源的列定义转为 columns 条目。"""
    cols = []
    for fm in field_mappings.get("fields", []):
        src_col = (fm.get("sources") or {}).get(source_key)
        if not src_col:
            continue
        col = dict(src_col)
        col["field_key"] = fm["field_key"]
        if not col.get("excel_column"):
            col["excel_column"] = fm.get("ui_label") or fm["field_key"]
        cols.append(col)
    return cols


def apply_field_mappings_to_sources(sources, field_mappings):
    """合并 field_mappings 与 sources/*.json 中的 columns（映射优先，避免重复 field_key）。"""
    for src in sources:
        src_key = src.get("key")
        mapped = _mapping_columns_for_source(field_mappings, src_key)
        mapped_keys = {c["field_key"] for c in mapped}
        extra = [c for c in src.get("columns", []) if c.get("field_key") not in mapped_keys]
        src["columns"] = mapped + extra
    return sources


def locked_fields_from_mappings(field_mappings):
    """导入后不可编辑的网页字段（field_key 列表）。"""
    locked = []
    for fm in field_mappings.get("fields", []):
        if fm.get("lock_on_import") and fm.get("field_key"):
            locked.append(fm["field_key"])
    return locked


def load_master_import_config(page="registration"):
    """按页面目录加载拆分后的主数据导入配置。"""
    page_dir = os.path.join(MASTER_IMPORT_DIR, page)
    index_path = os.path.join(page_dir, "index.json")
    if not os.path.exists(index_path):
        raise ValueError(f"主数据导入配置不存在: {page}")

    with open(index_path, encoding="utf-8") as f:
        index = json.load(f)

    field_mappings = load_field_mappings(page_dir, index)

    sources = []
    for key in index.get("source_keys", []):
        src_path = os.path.join(page_dir, "sources", f"{key}.json")
        with open(src_path, encoding="utf-8") as f:
            sources.append(json.load(f))
    apply_field_mappings_to_sources(sources, field_mappings)

    rules_path = os.path.join(page_dir, "stage_rules.json")
    with open(rules_path, encoding="utf-8") as f:
        stage_rules = json.load(f)

    cfg = dict(index)
    cfg["sources"] = sources
    cfg["stage_rules"] = stage_rules
    cfg["field_mappings"] = field_mappings
    cfg["registration_locked_fields"] = (
        index.get("registration_locked_fields") or locked_fields_from_mappings(field_mappings)
    )
    return cfg


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _text_match(pattern, text):
    if not pattern or text is None:
        return False
    text = str(text).strip()
    if not text:
        return False
    if "*" in pattern or "?" in pattern:
        return (fnmatch.fnmatch(text, pattern)
                or fnmatch.fnmatch(text.lower(), pattern.lower()))
    return pattern == text or pattern.lower() == text.lower()


def _column_patterns(col_cfg):
    seen = set()
    for key in ("excel_column", "field_key"):
        val = col_cfg.get(key)
        if val and val not in seen:
            seen.add(val)
            yield val
    for val in col_cfg.get("excel_aliases", []):
        if val and val not in seen:
            seen.add(val)
            yield val
    for val in col_cfg.get("match_patterns", []):
        if val and val not in seen:
            seen.add(val)
            yield val


def match_column(col_cfg, header):
    if not header:
        return False
    for pattern in _column_patterns(col_cfg):
        if _text_match(pattern, header):
            return True
    return False


def parse_master_header(source_cfg, header_row):
    """表头行 -> {col_index: field_key}"""
    headers = [_cell_str(h) for h in header_row]
    col_map = {}
    for idx, h in enumerate(headers):
        if not h:
            continue
        for col in source_cfg["columns"]:
            if match_column(col, h):
                col_map[idx] = col["field_key"]
                break
    required = [c["field_key"] for c in source_cfg["columns"] if c.get("required")]
    for rk in required:
        if rk not in col_map.values():
            label = next(c["excel_column"] for c in source_cfg["columns"] if c["field_key"] == rk)
            raise ValueError(f"Excel 中未找到必填列「{label}」")
    return col_map


def row_to_data(source_cfg, col_map, raw_row):
    data = {}
    for idx, key in col_map.items():
        if idx < len(raw_row):
            data[key] = _cell_str(raw_row[idx])
    return data


def _value_matches(expected, actual):
    actual = str(actual or "").strip()
    if expected == "*":
        return bool(actual)
    if "*" in expected or "?" in expected:
        return (fnmatch.fnmatch(actual, expected)
                or fnmatch.fnmatch(actual.lower(), expected.lower()))
    return actual == expected


def _rule_matches(data, rule):
    fields = rule.get("fields") or []
    values = rule.get("values") or []
    if len(fields) != len(values):
        return False
    for fk, expected in zip(fields, values):
        actual = data.get(fk, "")
        if not _value_matches(expected, actual):
            return False
    return True


def compute_current_stage(data, cfg=None):
    """根据 stage_rules 计算候选人当前流程阶段。"""
    cfg = cfg or load_master_import_config()
    field_key = cfg.get("current_stage_field", "current_stage")
    rules = sorted(cfg.get("stage_rules", {}).get("rules", []),
                   key=lambda r: r.get("priority", 0), reverse=True)
    for rule in rules:
        if _rule_matches(data, rule):
            stage = rule["stage"]
            data[field_key] = stage
            return stage
    data[field_key] = "registration"
    return "registration"


def stage_label_map():
    return {s["key"]: s.get("short_label") or s["label"] for s in load_stages_meta()}


def merge_candidate_data(old, incoming, overwrite_empty_only=False):
    """合并导入数据到已有候选人 JSON。"""
    merged = dict(old)
    for k, v in incoming.items():
        if k.startswith("_"):
            continue
        if not v:
            continue
        if overwrite_empty_only:
            if not str(merged.get(k, "") or "").strip():
                merged[k] = v
        elif str(merged.get(k, "") or "") != v:
            merged[k] = v
    return merged


def registration_locked_fields(cfg=None):
    cfg = cfg or load_master_import_config()
    return list(cfg.get("registration_locked_fields") or
                locked_fields_from_mappings(cfg.get("field_mappings") or {}))


def _select_workbook_sheet(wb, source_cfg, source_label=""):
    label = source_label or source_cfg.get("label") or source_cfg.get("key") or "Excel"
    name = (source_cfg.get("sheet_name") or "").strip()
    if name:
        if name not in wb.sheetnames:
            raise ValueError(
                f"「{label}」中未找到工作表「{name}」，当前工作表：{', '.join(wb.sheetnames)}"
            )
        return wb[name]
    idx = int(source_cfg.get("sheet_index", 0))
    if idx < 0 or idx >= len(wb.sheetnames):
        raise ValueError(
            f"「{label}」工作表索引 {idx} 无效，当前共 {len(wb.sheetnames)} 个工作表"
        )
    return wb[wb.sheetnames[idx]]


def parse_excel_file(path, source_cfg):
    """读取 Excel 指定工作表并解析为行数据列表（列映射见 sources + field_mappings.json）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    sheet = _select_workbook_sheet(wb, source_cfg)
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    col_map = parse_master_header(source_cfg, rows[0])
    return [row_to_data(source_cfg, col_map, raw) for raw in rows[1:]]


def merge_master_import_data(old, incoming, cfg=None):
    """主数据表合并：登记五字段以主表为准并锁定；其余字段常规覆盖合并。"""
    cfg = cfg or load_master_import_config()
    reg_locked = registration_locked_fields(cfg)
    merged = dict(old)
    locked = list(old.get("_master_locked_fields") or [])

    for k in reg_locked:
        val = incoming.get(k)
        if val and str(val).strip():
            merged[k] = str(val).strip()
            if k not in locked:
                locked.append(k)

    for k, v in incoming.items():
        if k.startswith("_") or k in reg_locked:
            continue
        if not v:
            continue
        nv = str(v).strip() if not isinstance(v, str) else v.strip()
        if not nv:
            continue
        if str(merged.get(k, "") or "") != nv:
            merged[k] = nv

    if locked:
        merged["_master_locked_fields"] = locked
        merged["_master_imported"] = True
    return merged


def _strip_internal_fields(data):
    return {k: v for k, v in data.items() if not k.startswith("_")}


def resolve_group_id(data, cfg, groups_list):
    """从行数据解析新建候选人所属分组。"""
    gm = cfg.get("group_mapping") or {}
    field_key = gm.get("field_key", "_group_name")
    raw = str(data.get(field_key, "") or "").strip()
    name_to_id = {g["name"]: g["id"] for g in groups_list}

    if raw:
        if raw in name_to_id:
            return name_to_id[raw]
        for g in groups_list:
            if _text_match(raw, g["name"]):
                return g["id"]
        for g in groups_list:
            for pat in gm.get("match_patterns", []):
                if _text_match(pat, g["name"]) and (_text_match(pat, raw) or raw in g["name"]):
                    return g["id"]

    fallback = gm.get("fallback", "first")
    if fallback == "first" and groups_list:
        return groups_list[0]["id"]
    return None


def build_global_candidate_index(db):
    """全局索引：简历编号 / 电话 / 姓名 -> 候选人记录（跨分组）。"""
    by_resume_id, by_phone, by_name = {}, {}, {}
    for row in db.execute("SELECT * FROM candidates").fetchall():
        data = json.loads(row["data"])
        if data.get("resume_id") and data["resume_id"] not in by_resume_id:
            by_resume_id[data["resume_id"]] = row
        if data.get("phone") and data["phone"] not in by_phone:
            by_phone[data["phone"]] = row
        if data.get("name") and data["name"] not in by_name:
            by_name[data["name"]] = row
    return by_resume_id, by_phone, by_name


def join_master_rows(app_rows, mgmt_rows, join_key="resume_id"):
    """按 join_key 合并 Application 与候选人管理两张表。"""
    mgmt_by_key = {}
    for row in mgmt_rows:
        key = str(row.get(join_key, "") or "").strip()
        if key:
            mgmt_by_key[key] = row

    merged, seen = [], set()
    for app_row in app_rows:
        key = str(app_row.get(join_key, "") or "").strip()
        if not key:
            continue
        combined = dict(app_row)
        if key in mgmt_by_key:
            combined = merge_candidate_data(combined, mgmt_by_key[key])
        merged.append(combined)
        seen.add(key)

    for key, mgmt_row in mgmt_by_key.items():
        if key in seen:
            continue
        if mgmt_row.get("name") or mgmt_row.get("phone"):
            merged.append(dict(mgmt_row))
    return merged


def run_dual_master_refresh(db, cfg, can_edit_fn, user):
    """从已上传的双表文件读取、关联、全局刷新全部候选人。"""
    from campus.services.master_import_store import both_files_ready, get_stored_files

    page = cfg.get("page", "registration")
    if not both_files_ready(page, cfg):
        raise ValueError("请先上传 Application*.xlsx 与 候选人管理*.xlsx 两个文件")

    files = get_stored_files(page, cfg)
    join_key = cfg.get("join_key", "resume_id")
    source_map = {s["key"]: s for s in cfg["sources"]}

    app_rows = parse_excel_file(files["application"]["path"], source_map["application"])
    mgmt_rows = parse_excel_file(files["candidate_mgmt"]["path"], source_map["candidate_mgmt"])
    rows_data = join_master_rows(app_rows, mgmt_rows, join_key)

    created, updated, skipped = apply_master_rows(
        rows_data, db, cfg, can_edit_fn, user)

    from campus.services.master_import_store import load_meta, save_meta
    meta = load_meta(page)
    meta["last_refresh"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta["last_refresh_stats"] = {
        "created": created, "updated": updated, "skipped": skipped,
        "application_rows": len(app_rows), "candidate_mgmt_rows": len(mgmt_rows),
        "merged_rows": len(rows_data),
    }
    save_meta(page, meta)
    return created, updated, skipped, meta["last_refresh_stats"]


def apply_master_rows(rows_data, db, cfg, can_edit_fn, user):
    """全局主表导入：跨记录匹配，按权限更新/新建。"""
    created = updated = skipped = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    by_resume_id, by_phone, by_name = build_global_candidate_index(db)
    match_keys = cfg.get("match_keys") or ["resume_id", "phone", "name"]

    for data in rows_data:
        if not data.get("name") and not data.get("resume_id"):
            skipped += 1
            continue

        compute_current_stage(data, cfg)
        match = None
        for mk in match_keys:
            val = str(data.get(mk, "") or "").strip()
            if not val:
                continue
            if mk == "resume_id":
                match = by_resume_id.get(val)
            elif mk == "phone":
                match = by_phone.get(val)
            elif mk == "name":
                match = by_name.get(val)
            if match:
                break

        if match:
            if not can_edit_fn(user):
                skipped += 1
                continue
            old = json.loads(match["data"])
            merged = merge_master_import_data(old, data, cfg)
            compute_current_stage(merged, cfg)
            if merged != old:
                db.execute(
                    "UPDATE candidates SET data=?, updated_at=? WHERE id=?",
                    (json.dumps(merged, ensure_ascii=False), now, match["id"]),
                )
                updated += 1
                match = db.execute("SELECT * FROM candidates WHERE id=?", (match["id"],)).fetchone()
                d = json.loads(match["data"])
                if d.get("resume_id"):
                    by_resume_id[d["resume_id"]] = match
                if d.get("phone"):
                    by_phone[d["phone"]] = match
                if d.get("name"):
                    by_name[d["name"]] = match
        else:
            if not data.get("name"):
                skipped += 1
                continue
            if not can_edit_fn(user):
                skipped += 1
                continue
            payload = _strip_internal_fields(data)
            payload = merge_master_import_data({}, data, cfg)
            cur = db.execute(
                "INSERT INTO candidates (group_id, data, created_at, updated_at) VALUES (?,?,?,?)",
                (None, json.dumps(payload, ensure_ascii=False), now, now),
            )
            row = db.execute("SELECT * FROM candidates WHERE id=?", (cur.lastrowid,)).fetchone()
            d = json.loads(row["data"])
            if d.get("resume_id"):
                by_resume_id[d["resume_id"]] = row
            if d.get("phone"):
                by_phone[d["phone"]] = row
            if d.get("name"):
                by_name[d["name"]] = row
            created += 1
    return created, updated, skipped
