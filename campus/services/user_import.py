# -*- coding: utf-8 -*-
"""用户批量导入/导出（权限管理）。"""
import io
import sqlite3

from openpyxl import Workbook, load_workbook

from campus.core.roles_store import role_keys, roles_payload
from campus.services.users import (
    apply_role_to_user,
    default_password_for_username,
    parse_user_profile_body,
    persist_user_columns,
    reserved_accounts,
    sync_registration_employee_snapshots,
    user_dict,
)

# (字段key, 表头中文)
USER_IMPORT_COLUMNS = [
    ("username", "工号"),
    ("display_name", "姓名"),
    ("role", "角色"),
    ("dept_level2", "二层部门"),
    ("dept_level3", "三层部门"),
    ("pl_group", "PL"),
    ("password", "密码"),
]


def _role_resolve_map():
    """角色 key / 中文名 → role key。"""
    m = {}
    for r in roles_payload():
        m[r["key"]] = r["key"]
        m[r["label"]] = r["key"]
    return m


def build_user_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "用户导入"
    for i, (_, label) in enumerate(USER_IMPORT_COLUMNS, start=1):
        ws.cell(row=1, column=i, value=label)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(label) * 2 + 2)
    ws.cell(row=2, column=1, value="（示例，导入前请删除本行）")
    ws.cell(row=2, column=2, value="张三")
    ws.cell(row=2, column=3, value="user")
    ws.cell(row=2, column=7, value="留空则默认同工号")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_users_workbook(db):
    wb = Workbook()
    ws = wb.active
    ws.title = "用户列表"
    headers = [label for _, label in USER_IMPORT_COLUMNS]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    role_labels = {r["key"]: r["label"] for r in roles_payload()}
    for ri, row in enumerate(db.execute("SELECT * FROM users ORDER BY id").fetchall(), start=2):
        u = user_dict(row)
        ws.cell(row=ri, column=1, value=u["username"])
        ws.cell(row=ri, column=2, value=u.get("display_name") or "")
        ws.cell(row=ri, column=3, value=role_labels.get(u.get("role"), u.get("role") or ""))
        ws.cell(row=ri, column=4, value=u.get("dept_level2") or "")
        ws.cell(row=ri, column=5, value=u.get("dept_level3") or "")
        ws.cell(row=ri, column=6, value=u.get("pl_group") or "")
        ws.cell(row=ri, column=7, value="")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _header_index_map(header_row):
    label_to_key = {label: key for key, label in USER_IMPORT_COLUMNS}
    idx = {}
    for col, cell in enumerate(header_row, start=1):
        val = str(cell.value or "").strip()
        if val in label_to_key:
            idx[label_to_key[val]] = col
    return idx


def import_users_from_workbook(db, file_storage, *, update_existing=True):
    """解析 Excel 批量创建/更新用户。返回统计 dict。"""
    try:
        wb = load_workbook(file_storage, data_only=True)
    except Exception as exc:
        raise ValueError("文件解析失败，请上传 .xlsx 格式文件") from exc
    ws = wb.active
    rows = list(ws.iter_rows())
    if not rows:
        raise ValueError("表格为空")
    col_map = _header_index_map(rows[0])
    if "username" not in col_map or "display_name" not in col_map:
        raise ValueError("表头须包含「工号」「姓名」列，请使用系统提供的导入模板")

    role_map = _role_resolve_map()
    valid_roles = set(role_keys())
    stats = {"created": 0, "updated": 0, "skipped": 0, "errors": []}

    def cell_val(row, key):
        c = col_map.get(key)
        if not c or c > len(row):
            return ""
        return str(row[c - 1].value or "").strip()

    for line_no, row in enumerate(rows[1:], start=2):
        username = cell_val(row, "username")
        if not username or username.startswith("（"):
            continue
        display_name = cell_val(row, "display_name")
        if not display_name:
            stats["errors"].append(f"第{line_no}行：姓名不能为空")
            stats["skipped"] += 1
            continue
        if username in reserved_accounts() and username != "admin":
            stats["errors"].append(f"第{line_no}行：工号「{username}」为保留账号，不可导入")
            stats["skipped"] += 1
            continue

        role_raw = cell_val(row, "role") or "user"
        role = role_map.get(role_raw, role_raw)
        if role not in valid_roles:
            stats["errors"].append(f"第{line_no}行：角色「{role_raw}」不合法")
            stats["skipped"] += 1
            continue

        body = {
            "username": username,
            "display_name": display_name,
            "role": role,
            "dept_level2": cell_val(row, "dept_level2"),
            "dept_level3": cell_val(row, "dept_level3"),
            "pl_group": cell_val(row, "pl_group"),
        }
        err, fields = parse_user_profile_body(body)
        if err:
            stats["errors"].append(f"第{line_no}行：{err}")
            stats["skipped"] += 1
            continue

        pwd = cell_val(row, "password")
        existing = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        try:
            if existing:
                if not update_existing:
                    stats["skipped"] += 1
                    continue
                persist_user_columns(db, existing["id"], fields, role,
                                     password=pwd or None)
                apply_role_to_user(db, existing["id"], role)
                sync_registration_employee_snapshots(db, username)
                stats["updated"] += 1
            else:
                password = pwd or default_password_for_username(username)
                persist_user_columns(db, None, fields, role, password=password,
                                     is_create=True, username=username)
                uid = db.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()["id"]
                apply_role_to_user(db, uid, role)
                stats["created"] += 1
        except sqlite3.IntegrityError:
            stats["errors"].append(f"第{line_no}行：工号「{username}」已存在")
            stats["skipped"] += 1
        except Exception as e:
            stats["errors"].append(f"第{line_no}行：{e}")
            stats["skipped"] += 1

    return stats
