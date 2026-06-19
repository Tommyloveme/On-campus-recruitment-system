# -*- coding: utf-8 -*-
"""Excel 导入/导出与主数据表导入。"""
import io
import json
from datetime import datetime

from flask import Blueprint, g, jsonify, request, send_file
from openpyxl import Workbook, load_workbook

from campus.auth.decorators import login_required
from campus.auth.permissions import GLOBAL_VIEW_ROLES, can_edit_group, can_view_group
from campus.config_loader import (
    get_stage_meta,
    importable_fields,
    load_stage_fields,
    match_import_header,
    validate_stage,
)
from campus.db.connection import get_db, now_str
from campus.logging_util import log, who
from campus.services.audit import add_log
from campus.services.candidates import group_name_map
from campus.stage_engine import (
    apply_master_rows,
    compute_current_stage,
    load_master_import_config,
    parse_master_header,
    row_to_data,
)

bp = Blueprint("import_export", __name__)


@bp.get("/api/import/template")
@login_required
def api_template():
    stage = request.args.get("stage", "registration")
    try:
        validate_stage(stage)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    group_id = None if g.user["role"] in GLOBAL_VIEW_ROLES else g.user["group_id"]
    fields = importable_fields(stage, group_id)
    stage_label = get_stage_meta(stage)["label"]
    wb = Workbook()
    ws = wb.active
    ws.title = stage_label
    for i, f in enumerate(fields, start=1):
        ws.cell(row=1, column=i, value=f["excel_column"])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(f["excel_column"]) * 2 + 4)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{stage_label}导入模板.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/api/import")
@login_required
def api_import():
    if "file" not in request.files:
        return jsonify({"error": "请选择Excel文件"}), 400
    stage = request.form.get("stage", "registration")
    try:
        validate_stage(stage)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    group_id = request.form.get("group_id", type=int) or g.user["group_id"]
    if not group_id:
        return jsonify({"error": "请指定导入的分组"}), 400
    if not can_edit_group(g.user, group_id):
        log.warning("Excel导入权限拒绝 %s group_id=%s stage=%s", who(g.user), group_id, stage)
        return jsonify({"error": "无该分组的编辑权限"}), 403

    try:
        wb = load_workbook(request.files["file"], data_only=True)
    except Exception:
        return jsonify({"error": "文件解析失败，请上传 .xlsx 格式文件"}), 400
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "Excel内容为空"}), 400

    fields = importable_fields(stage)
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_map = {}
    for idx, h in enumerate(header):
        for f in fields:
            if match_import_header(f, h):
                col_map[idx] = f["key"]
                break
    if "name" not in col_map.values():
        return jsonify({"error": "Excel中未找到“候选人”列，请参考导入模板"}), 400
    log.debug("Excel导入列映射 %s group_id=%s stage=%s cols=%s", who(g.user), group_id, stage, col_map)

    def cell_str(v):
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d")
        return str(v).strip()

    db = get_db()
    existing = db.execute("SELECT * FROM candidates WHERE group_id=?", (group_id,)).fetchall()
    by_phone = {}
    by_name = {}
    for r in existing:
        d = json.loads(r["data"])
        if d.get("phone"):
            by_phone[d["phone"]] = r
        if d.get("name"):
            by_name[d["name"]] = r

    created = updated = skipped = 0
    for raw in rows[1:]:
        data = {}
        for idx, key in col_map.items():
            if idx < len(raw):
                data[key] = cell_str(raw[idx])
        if not data.get("name"):
            skipped += 1
            continue
        compute_current_stage(data)
        match = by_phone.get(data.get("phone", "")) or by_name.get(data["name"])
        if match:
            old = json.loads(match["data"])
            merged = dict(old)
            changed = False
            for k, v in data.items():
                if v and str(old.get(k, "")) != v:
                    merged[k] = v
                    changed = True
            if changed:
                compute_current_stage(merged)
                db.execute("UPDATE candidates SET data=?, updated_at=? WHERE id=?",
                           (json.dumps(merged, ensure_ascii=False), now_str(), match["id"]))
                updated += 1
        else:
            cur = db.execute("INSERT INTO candidates (group_id, data, created_at, updated_at) VALUES (?,?,?,?)",
                             (group_id, json.dumps(data, ensure_ascii=False), now_str(), now_str()))
            row = db.execute("SELECT * FROM candidates WHERE id=?", (cur.lastrowid,)).fetchone()
            by_name[data["name"]] = row
            if data.get("phone"):
                by_phone[data["phone"]] = row
            created += 1

    gname = group_name_map().get(group_id, "")
    stage_label = get_stage_meta(stage)["label"]
    add_log(g.user, "import",
            f"{g.user['display_name']} 通过Excel向「{gname}」{stage_label}导入：新增{created}人，更新{updated}人"
            + (f"，跳过{skipped}行" if skipped else ""),
            None, None, group_id)
    db.commit()
    log.info("Excel导入 %s group=%s stage=%s 新增%d 更新%d 跳过%d",
             who(g.user), gname, stage, created, updated, skipped)
    return jsonify({"ok": True, "created": created, "updated": updated, "skipped": skipped})


@bp.get("/api/master-import/config")
@login_required
def api_master_import_config():
    page = request.args.get("page", "registration")
    try:
        cfg = load_master_import_config(page)
    except ValueError as e:
        return jsonify({"error": str(e)}), 404
    return jsonify({
        "page": cfg.get("page", page),
        "sources": [{"key": s["key"], "label": s["label"], "description": s.get("description", "")}
                    for s in cfg.get("sources", [])],
        "stage_rules": cfg.get("stage_rules", {}),
        "current_stage_field": cfg.get("current_stage_field", "current_stage"),
        "global_import": bool(cfg.get("global_import", True)),
    })


@bp.post("/api/master-import")
@login_required
def api_master_import():
    if "file" not in request.files:
        return jsonify({"error": "请选择Excel文件"}), 400
    source_key = request.form.get("source", "primary")
    page = request.form.get("page", "registration")

    try:
        cfg = load_master_import_config(page)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    source = next((s for s in cfg.get("sources", []) if s["key"] == source_key), None)
    if not source:
        return jsonify({"error": f"未知数据源: {source_key}"}), 400

    try:
        wb = load_workbook(request.files["file"], data_only=True)
    except Exception:
        return jsonify({"error": "文件解析失败，请上传 .xlsx 格式文件"}), 400
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "Excel内容为空"}), 400

    try:
        col_map = parse_master_header(source, rows[0])
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    rows_data = [row_to_data(source, col_map, raw) for raw in rows[1:]]

    db = get_db()
    groups_list = [{"id": r["id"], "name": r["name"]}
                   for r in db.execute("SELECT id, name FROM groups ORDER BY id").fetchall()]
    if not groups_list:
        return jsonify({"error": "系统中暂无分组，无法导入"}), 400

    created, updated, skipped = apply_master_rows(
        rows_data, db, cfg, can_edit_group, g.user, groups_list)
    add_log(g.user, "import",
            f"{g.user['display_name']} 通过主数据表「{source['label']}」全局导入："
            f"新增{created}人，更新{updated}人" + (f"，跳过{skipped}行" if skipped else ""))
    db.commit()
    log.info("主数据导入 %s source=%s page=%s +%d ~%d skip%d",
             who(g.user), source_key, page, created, updated, skipped)
    return jsonify({"ok": True, "created": created, "updated": updated, "skipped": skipped})


@bp.post("/api/candidates/export")
@login_required
def api_candidates_export():
    b = request.get_json(force=True)
    ids = b.get("ids") or []
    stage = b.get("stage", "registration")
    try:
        validate_stage(stage)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    if not ids:
        return jsonify({"error": "请先勾选候选人"}), 400
    db = get_db()
    placeholders = ",".join("?" * len(ids))
    rows = db.execute(f"SELECT * FROM candidates WHERE id IN ({placeholders})", ids).fetchall()

    group_id = None if g.user["role"] in GLOBAL_VIEW_ROLES else g.user["group_id"]
    fields = [f for f in load_stage_fields(stage, group_id) if f.get("visible", True)]
    names = group_name_map()
    stage_label = get_stage_meta(stage)["label"]

    wb = Workbook()
    ws = wb.active
    ws.title = stage_label
    headers = ["二层部门"] + [f["label"] for f in fields]
    ws.append(headers)
    exported = 0
    for row in rows:
        if not can_view_group(g.user, row["group_id"]):
            continue
        data = json.loads(row["data"])
        ws.append([names.get(row["group_id"], "")] + [data.get(f["key"], "") for f in fields])
        exported += 1
    if exported == 0:
        return jsonify({"error": "选中的候选人均无导出权限"}), 403
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(h) * 2 + 4)

    add_log(g.user, "export",
            f"{g.user['display_name']} 从{stage_label}导出了 {exported} 名候选人的Excel数据")
    db.commit()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = send_file(buf, as_attachment=True, download_name=f"{stage_label}导出_{ts}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["X-Export-Count"] = str(exported)
    return resp
